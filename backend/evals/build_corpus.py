"""Render the synthetic eval corpus to evals/corpus/ (deterministic; outputs are gitignored).

    python -m evals.build_corpus
"""
from pathlib import Path

from evals.corpus_data import HEADCOUNT, PRODUCTS, RESUMES, REVENUE, SKUS, render_product_spec, render_resume

CORPUS_DIR = Path(__file__).parent / "corpus"

# Resumes rendered as .docx vs .pdf to exercise both parsers.
_DOCX_RESUMES = {"resume_marcus_bell"}


def corpus_texts() -> dict[str, str]:
    """filename -> plain text, used to prove gold evidence exists in exactly one document."""
    texts: dict[str, str] = {}
    for p in PRODUCTS:
        texts[f"spec_{p['slug']}.md"] = render_product_spec(p)
    for r in RESUMES:
        ext = "docx" if r["slug"] in _DOCX_RESUMES else "pdf"
        texts[f"{r['slug']}.{ext}"] = render_resume(r)
    q3 = " ".join(" ".join(str(c) for c in row) for row in SKUS)
    q2 = " ".join(f"{sku} {name} {cat} {q2c} {moq}" for sku, name, cat, _q3, q2c, moq in SKUS)
    texts["pricing.xlsx"] = f"Q3 Pricing {q3} Q2 Pricing {q2}"
    rev = " ".join(" ".join(str(c) for c in row) for row in REVENUE)
    head = " ".join(" ".join(str(c) for c in row) for row in HEADCOUNT)
    texts["financials_q3.xlsx"] = f"Revenue {rev} Headcount {head}"
    return texts


def build() -> list[Path]:
    import fitz
    from docx import Document
    from openpyxl import Workbook

    CORPUS_DIR.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []

    for p in PRODUCTS:
        path = CORPUS_DIR / f"spec_{p['slug']}.md"
        path.write_text(render_product_spec(p), encoding="utf-8")
        paths.append(path)

    for r in RESUMES:
        text = render_resume(r)
        if r["slug"] in _DOCX_RESUMES:
            path = CORPUS_DIR / f"{r['slug']}.docx"
            doc = Document()
            for line in text.splitlines():
                doc.add_paragraph(line)
            doc.save(path)
        else:
            path = CORPUS_DIR / f"{r['slug']}.pdf"
            pdf = fitz.open()
            page = pdf.new_page()
            page.insert_textbox(fitz.Rect(50, 50, 545, 790), text, fontsize=10)
            pdf.save(path)
            pdf.close()
        paths.append(path)

    wb = Workbook()
    q3 = wb.active
    q3.title = "Q3 Pricing"
    q3.append(["SKU", "Product", "Category", "Unit Cost (USD)", "MOQ"])
    q2 = wb.create_sheet("Q2 Pricing")
    q2.append(["SKU", "Product", "Category", "Unit Cost (USD)", "MOQ"])
    for sku, name, cat, q3c, q2c, moq in SKUS:
        q3.append([sku, name, cat, q3c, moq])
        q2.append([sku, name, cat, q2c, moq])
    path = CORPUS_DIR / "pricing.xlsx"
    wb.save(path)
    paths.append(path)

    wb = Workbook()
    rev = wb.active
    rev.title = "Revenue"
    rev.append(["Region", "Q1 (USD k)", "Q2 (USD k)", "Q3 (USD k)"])
    for row in REVENUE:
        rev.append(list(row))
    head = wb.create_sheet("Headcount")
    head.append(["Department", "Headcount"])
    for row in HEADCOUNT:
        head.append(list(row))
    path = CORPUS_DIR / "financials_q3.xlsx"
    wb.save(path)
    paths.append(path)

    return paths


if __name__ == "__main__":
    for p in build():
        print(p.name)
